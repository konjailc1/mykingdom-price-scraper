"""
mykingdom-price-monitor
=======================
A beginner-friendly script that reads the PUBLIC product list of the toy store
https://www.mykingdom.com.vn so you can monitor toy prices and promotions.

The store runs on Shopify. Every Shopify store offers a public "products.json"
page that lists products as structured data (JSON), so we do NOT need to parse
messy HTML. We simply ask for page 1, page 2, ... until a page comes back empty.

What this script does:
  1. Downloads /products.json?limit=250&page=N politely.
  2. Turns every product VARIANT (e.g. a size or colour) into one table row.
  3. Saves the table to products.csv and a nicely formatted products.xlsx.
  4. Prints a summary (in-stock %, average price, top brands, biggest discounts).

Example usage:
    python main.py                       # 3 pages -> sample_output/
    python main.py --max-pages 2         # only 2 pages
    python main.py --max-pages 10 --out data
"""

import argparse
import sys
import time
from pathlib import Path
from urllib.parse import urlparse

import pandas as pd
import requests
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ----------------------------------------------------------------------------
# Settings (constants). Change them here if you want different behaviour.
# ----------------------------------------------------------------------------

# The ONLY website this script is allowed to talk to.
BASE_URL = "https://www.mykingdom.com.vn"
ALLOWED_HOST = "www.mykingdom.com.vn"

# The ONLY page we download. It is allowed by the site's robots.txt.
PRODUCTS_PATH = "/products.json"

# robots.txt asks robots to stay away from these private areas. Our code
# double-checks every URL and refuses to touch them.
FORBIDDEN_PREFIXES = ("/cart", "/checkout", "/account", "/admin", "/orders")

# A custom User-Agent tells the website who is visiting. Being honest is polite.
USER_AGENT = (
    "mykingdom-price-monitor/1.0 "
    "(learning project; public product data only; reads /products.json)"
)

REQUEST_TIMEOUT = 15   # seconds to wait for a response before giving up
MAX_RETRIES = 3        # how many times to try a page before giving up
DELAY_SECONDS = 1.5    # pause between requests so we don't hammer the server
PAGE_SIZE = 250        # Shopify's maximum number of products per page
DEFAULT_MAX_PAGES = 3  # default number of pages to download

# The columns of our table, in the order they will appear in the files.
COLUMNS = [
    "product_id", "title", "vendor", "product_type", "tags", "sku",
    "price_vnd", "compare_at_price_vnd", "discount_pct", "available",
    "product_url", "image_url", "created_at",
]


# ----------------------------------------------------------------------------
# Step 1: downloading pages
# ----------------------------------------------------------------------------

def check_url_is_allowed(url):
    """Stop the program if a URL is outside what we promised to scrape."""
    parsed = urlparse(url)
    if parsed.hostname != ALLOWED_HOST:
        raise ValueError(f"Refusing to scrape other sites: {url}")
    if parsed.path.startswith(FORBIDDEN_PREFIXES):
        raise ValueError(f"Refusing to touch a private area (robots.txt): {url}")
    if parsed.path != PRODUCTS_PATH:
        raise ValueError(f"This script only reads {PRODUCTS_PATH}: {url}")


def fetch_products_page(session, page):
    """Download one page of products and return the list of products (may be empty).

    Tries up to MAX_RETRIES times. If every attempt fails, raises the last error.
    """
    url = BASE_URL + PRODUCTS_PATH
    check_url_is_allowed(url)
    params = {"limit": PAGE_SIZE, "page": page}

    last_error = None
    for attempt in range(1, MAX_RETRIES + 1):
        try:
            response = session.get(url, params=params, timeout=REQUEST_TIMEOUT)

            # HTTP 429 means "you are going too fast". The site may tell us how
            # long to wait in the Retry-After header, so we respect it.
            if response.status_code == 429:
                wait = _retry_after_seconds(response, default=DELAY_SECONDS * attempt * 4)
                print(f"  Server asked us to slow down. Waiting {wait:.0f}s...")
                time.sleep(wait)
                raise requests.HTTPError("429 Too Many Requests", response=response)

            # raise_for_status() turns HTTP errors (404, 500, ...) into exceptions.
            response.raise_for_status()
            # .json() converts the JSON text into Python dictionaries and lists.
            return response.json().get("products", [])
        except (requests.RequestException, ValueError) as error:
            last_error = error
            print(f"  Attempt {attempt}/{MAX_RETRIES} failed for page {page}: {error}")
            if attempt < MAX_RETRIES:
                time.sleep(DELAY_SECONDS * attempt)  # wait a bit longer each time

    raise last_error


def _retry_after_seconds(response, default):
    """Read the Retry-After header (in seconds); fall back to `default`."""
    try:
        return min(float(response.headers.get("Retry-After", default)), 60)
    except ValueError:
        return default


# ----------------------------------------------------------------------------
# Step 2: turning products into table rows
# ----------------------------------------------------------------------------

def to_number(value):
    """Convert Shopify's text price like '559000' to the number 559000.0 (or None)."""
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def calc_discount_pct(price, compare_at_price):
    """Discount in percent. Example: price 80, was 100 -> 20.0.

    Shopify's 'compare_at_price' is the original ("crossed out") price. If it is
    missing or not higher than the price, there is no discount, so we return 0.
    """
    if price is None or compare_at_price is None or compare_at_price <= price:
        return 0.0
    return round((compare_at_price - price) / compare_at_price * 100, 1)


def products_to_rows(products):
    """Create one row (dictionary) per variant of every product."""
    rows = []
    for product in products:
        # The first image of the product (some products have no images).
        images = product.get("images") or []
        image_url = images[0].get("src") if images else None

        product_url = f"{BASE_URL}/products/{product.get('handle', '')}"
        # Tags arrive as a list; join them into one text so it fits one cell.
        tags = ", ".join(product.get("tags") or [])

        for variant in product.get("variants", []):
            price = to_number(variant.get("price"))
            compare_at = to_number(variant.get("compare_at_price"))
            rows.append({
                "product_id": product.get("id"),
                "title": product.get("title"),
                "vendor": product.get("vendor"),
                "product_type": product.get("product_type"),
                "tags": tags,
                "sku": variant.get("sku"),
                "price_vnd": price,
                "compare_at_price_vnd": compare_at,
                "discount_pct": calc_discount_pct(price, compare_at),
                "available": bool(variant.get("available")),
                "product_url": product_url,
                "image_url": image_url,
                "created_at": product.get("created_at"),
            })
    return rows


# ----------------------------------------------------------------------------
# Step 3: crawling (looping over pages)
# ----------------------------------------------------------------------------

def crawl(max_pages):
    """Download up to `max_pages` pages and return all rows found."""
    # A Session reuses the connection and sends the same headers every time.
    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT})

    all_rows = []
    for page in range(1, max_pages + 1):
        print(f"Downloading page {page}/{max_pages} ...")

        try:
            products = fetch_products_page(session, page)
        except (requests.RequestException, ValueError):
            print("  Giving up on this page. Stopping and keeping what we have.")
            break

        if not products:  # an empty list means we reached the end of the catalogue
            print("  No more products. Reached the end of the catalogue.")
            break

        rows = products_to_rows(products)
        all_rows.extend(rows)
        print(f"  {len(products)} products, {len(rows)} variants")

        # Be polite: wait before the next request (only if there is one).
        if page < max_pages:
            time.sleep(DELAY_SECONDS)

    return all_rows


# ----------------------------------------------------------------------------
# Step 4: saving files
# ----------------------------------------------------------------------------

def save_excel(df, path):
    """Write the Excel file with a styled header, VND number format and auto-fit columns."""
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="products")
        sheet = writer.sheets["products"]

        # --- Formatted header: bold white text on a dark blue background ---
        header_fill = PatternFill("solid", fgColor="1F4E78")
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        sheet.freeze_panes = "A2"              # keep the header visible when scrolling
        sheet.auto_filter.ref = sheet.dimensions  # add filter arrows to the header

        # --- Number formats per column ---
        formats = {
            "product_id": "0",                  # avoid scientific notation (1.5E+13)
            "price_vnd": "#,##0",               # 559,000
            "compare_at_price_vnd": "#,##0",
            "discount_pct": "0.0",
        }
        for column_index, name in enumerate(df.columns, start=1):
            if name in formats:
                for row in range(2, sheet.max_row + 1):
                    sheet.cell(row=row, column=column_index).number_format = formats[name]

        # --- Autofit: set each column as wide as its longest value (max 60) ---
        for column_index, name in enumerate(df.columns, start=1):
            longest = max(
                [len(str(name))] + [len(str(v)) for v in df[name].head(1000).fillna("")]
            )
            # Numbers get thousands separators, so add a little extra room.
            sheet.column_dimensions[get_column_letter(column_index)].width = min(longest + 3, 60)


def save_results(df, out_dir):
    """Save the DataFrame as products.csv and products.xlsx inside `out_dir`."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)  # create the folder if missing

    csv_path = out_dir / "products.csv"
    xlsx_path = out_dir / "products.xlsx"

    # utf-8-sig adds a small marker so Excel shows Vietnamese letters correctly.
    df.to_csv(csv_path, index=False, encoding="utf-8-sig")
    save_excel(df, xlsx_path)

    print(f"\nSaved: {csv_path}")
    print(f"Saved: {xlsx_path}")


# ----------------------------------------------------------------------------
# Step 5: summary
# ----------------------------------------------------------------------------

def print_summary(df):
    """Print the most useful numbers about the scraped data."""
    total_products = df["product_id"].nunique()
    in_stock_pct = df["available"].mean() * 100          # share of variants in stock
    avg_price = df["price_vnd"].mean()

    print("\n===== SUMMARY =====")
    print(f"Total products   : {total_products} ({len(df)} variants)")
    print(f"In stock         : {in_stock_pct:.1f}% of variants")
    print(f"Average price    : {avg_price:,.0f} VND")

    # Top 10 brands: count each product once, even if it has many variants.
    print("\nTop 10 brands by number of products:")
    brands = df.drop_duplicates("product_id")["vendor"].value_counts().head(10)
    for position, (brand, count) in enumerate(brands.items(), start=1):
        print(f"  {position:>2}. {brand} ({count})")

    # Top 10 discounts: only rows that really have a discount.
    print("\nTop 10 biggest discounts:")
    discounted = df[df["discount_pct"] > 0].nlargest(10, "discount_pct")
    if discounted.empty:
        print("  No discounted products found in this data.")
    for position, (_, row) in enumerate(discounted.iterrows(), start=1):
        print(
            f"  {position:>2}. -{row['discount_pct']:.1f}%  "
            f"{row['price_vnd']:,.0f} (was {row['compare_at_price_vnd']:,.0f}) VND  "
            f"{row['title']} [SKU: {row['sku']}]"
        )


# ----------------------------------------------------------------------------
# Command-line interface
# ----------------------------------------------------------------------------

def positive_int(value):
    """argparse helper: accept only whole numbers >= 1."""
    number = int(value)
    if number < 1:
        raise argparse.ArgumentTypeError("must be 1 or greater")
    return number


def parse_args():
    parser = argparse.ArgumentParser(
        description="Collect public product prices from https://www.mykingdom.com.vn"
    )
    parser.add_argument(
        "--max-pages", type=positive_int, default=DEFAULT_MAX_PAGES,
        help=f"maximum number of pages (250 products each) to download (default: {DEFAULT_MAX_PAGES})",
    )
    parser.add_argument(
        "--out", default="sample_output",
        help="folder where products.csv and products.xlsx are saved (default: sample_output)",
    )
    return parser.parse_args()


def main():
    # Make sure Vietnamese letters can be printed in the Windows console.
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    args = parse_args()

    rows = crawl(args.max_pages)
    if not rows:
        print("No products were collected. Check your internet connection and try again.")
        sys.exit(1)

    # A pandas DataFrame is like a spreadsheet table in memory.
    df = pd.DataFrame(rows, columns=COLUMNS)

    save_results(df, args.out)
    print_summary(df)


# This line means: only run main() when the file is executed directly
# (python main.py), not when it is imported from another file.
if __name__ == "__main__":
    main()
